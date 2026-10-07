#!/usr/bin/env python3
"""
CC Medical Inventory - spreadsheet import.

Step 1 (review):  python3 tools/import_inventory.py review Inventory_Ultrasound.xlsx Import_Review.xlsx
    Reads the Location sheet and writes a review workbook: the original columns beside the
    fields the app will store, with rows that need a person marked CHECK.

Step 2 (sql):     python3 tools/import_inventory.py sql Import_Review.xlsx import.sql
    Reads the (corrected) review workbook and writes SQL for:
    npx wrangler d1 execute cc-inventory --remote --file=import.sql

Step 2, simpler (json): python3 tools/import_inventory.py json Import_Review.xlsx CC_Medical_Import.json
    Same import as a file an admin uploads in the app: Admin > Import inventory.
    Or (csv): python3 tools/import_inventory.py csv Import_Review.xlsx CC_Medical_Inventory.csv

Needs: Python 3.9+ and openpyxl (pip install openpyxl).
Re-running the SQL is safe: rows already imported are skipped.
"""
import datetime as dt
from decimal import Decimal
import hashlib
import json
import os
import re
import sys
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------------------------
# Rules
# --------------------------------------------------------------------------------------------
STOCK_STATUSES = ["In stock", "Pending", "On loan", "Out on rental", "Sold / shipped"]
CLOSED = {"Sold / shipped"}

# Words at the start of MODULE that name the maker.
PREFIX_MAKERS = [
    (r"ZONARE\s+MINDRAY", "Mindray / Zonare"), (r"ZONARE", "Mindray / Zonare"),
    (r"SIEMENS\s+ACUSON", "Siemens"), (r"SIEMENS", "Siemens"), (r"ACUSON", "Siemens"),
    (r"SON+O?SO?ITE|SONSITE|SONSOITE", "SonoSite / Fujifilm"),
    (r"SAMSUNG", "Samsung"), (r"BK", "BK Medical"), (r"HITACHI", "Hitachi"),
    (r"CANON", "Canon / Toshiba"), (r"PHILL?IPS", "Philips"),
]
# Maker inferred from the model when no maker is written. First match wins.
MODEL_MAKERS = [
    (r"^(PVT|PVI|PLT|PLI|PLU|PST|PVU|PVC)[- ]", "Canon / Toshiba"),
    (r"^(HFL|HSL|L25X|L38X|RC60X|RP19X|ICTX|C11X|C60X|L19-5|19-5|X-?PORTE|EDGE|SII$|SLL$)", "SonoSite / Fujifilm"),
    (r"(^|\s)(LOGIQ|VIVID|VOLUSON|LE10)\b|6208000|GA2000", "GE Healthcare"),
    (r"\b(COMPACT|EXP|CLEARVUE|ERGO|EMT|EPIQ|IE33|CX50|HARMONY|FUSION|MORPHEUS|RAFI)\b|4535\d{6,}|M1671A", "Philips"),
    (r"^(18L6|10L4|14L5|4C1|4V1|5C1|9C3|MC9-4|9EC4|Z6MS|CH5-2|P4-2|10EV3|TE-V5MS|SWIFTLINK|S2000|SC2000)", "Siemens"),
    (r"^(E9-4|L10-5|L8-3)$", "Mindray / Zonare"),
    (r"^(RIC|RAB|RM\d|RSP|IC\d|ML\d|M\dS|AB2|E8C|P2D|P6D|6TC|9T\b|9L\b|12L\b|6VT|4VC|3SC|3S-|8C-|4C-|10T|11L|12S|6S-|3CRF|C1-5|C1-6|C2-7|C2-9|C3-10|L2-9|L3-12|L4-12T|L4-20T|L6-12|L6-24|L8-18I|L10-22)", "GE Healthcare"),
    (r"-(D|RS)$", "GE Healthcare"),
    (r"^(C5-1|C6-2|C8-5|C9-2|C9-3V|C9-4V|C10-3V|C10-4EC|X5-1|X7-2T|X8-2T|X11-4T|XL14-3|S4-2|S5-1|S7-3T|S8-3|S9-2|S12-4|L12-3|L12-4|L12-5|L15-7IO|L18-5|L9-3|EL18-4|MC7-2|MC12-3|V6-2|V9-2|3D9-3V|D2CWC)", "Philips"),
]
SYSTEM_PART = r"\b(BOARD|POWER|SUPPLY|MONITOR|LCD|PANEL|BATTERY|TRACKBALL|MODULE|PCBA|TGC|ASSY|ASSEMBLY|COMPUTER|REGULATOR|TOUCHSCREEN|TOUCH|MECHANISM|CONSOLE|HUB|ACB|LOWER OP)\b"
ACCESSORY = r"\b(ECG|CABLES?|LEADS?|ADAPTOR|ADAPTER|DOCK?|SWIFTLINK|TRUNK|GRABBER)\b"
SYSTEM = r"^(CX50|LOGIQ E R7 VET|X-?PORTE ENGINE|SII|EDGE II)$"

# Facility names written in INFO -> one customer name. Spellings that differ are grouped here.
# The real list is CC Medical's customer list, so it lives in tools/customer_names.json, which
# stays off GitHub (.gitignore). Copy customer_names.example.json to start one.
def _load_customers():
    here = os.path.dirname(os.path.abspath(__file__))
    for name in ("customer_names.json", "customer_names.example.json"):
        f = os.path.join(here, name)
        if os.path.exists(f):
            with open(f) as fh:
                return [(e["match"], e["name"]) for e in json.load(fh)]
    return []
CUSTOMERS = _load_customers()
NOT_A_CUSTOMER = r"\b(DO NOT|DON.?T)\b"
CUSTOMER_STATUSES = {"On loan", "Sold / shipped", "Pending", "Out on rental"}


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def status_and_bin(loc):
    u = clean(loc).upper()
    if u in ("SOLD", "SHIPPED"):
        return "Sold / shipped", "", None
    if u in ("LOAN", "LOANER"):
        return "On loan", "", None
    if u == "PENDING":
        return "Pending", "", None
    if u == "RENTAL":
        return "Out on rental", "", None
    if re.fullmatch(r"[A-Z]\d(\s*/\s*[A-Z]\d)?", u):
        return "In stock", re.sub(r"\s*/\s*", "/", u), None
    if u in ("", "-"):
        return "In stock", "", "No location in the spreadsheet"
    return "In stock", "", f'Location "{clean(loc)}" is not a shelf code'


def parse_module(raw):
    """MODULE text -> (manufacturer, model, part_number, condition_hint, maker_guessed)."""
    t = clean(raw).upper()
    t = t.replace("RSNEW", "RS NEW").replace("18L6HD", "18L6 HD")
    maker = ""
    for pat, name in PREFIX_MAKERS:
        m = re.match(rf"^({pat})\b\s*", t)
        if m:
            maker, t = name, t[m.end():]
            break
    # Part numbers: Philips 4535..., or digits / codes inside brackets, or after P/N or PN:
    parts = re.findall(r"4535\d{6,}[A-Z]?", t)
    for inner in re.findall(r"\(([^()]*)\)?", t):
        w = inner.strip()
        if re.fullmatch(r"(P/?N:?\s*)?[A-Z]{0,4}\d[\w-]{4,}", w) and not re.fullmatch(r"NEW|REFURBISHED", w):
            parts.append(re.sub(r"^P/?N:?\s*", "", w))
    for w in re.findall(r"P/?N:?\s*([\w-]{5,})", t):
        parts.append(w)
    part = next((p for p in parts if p), "")
    cond = ""
    if re.search(r"\bREFURB(ISHED)?\b", t):
        cond = "Refurbished"
    elif re.search(r"\bNEW\b", t):
        cond = "New"
    # Remove condition words, part numbers and stray brackets from the model.
    t = re.sub(r"\(\s*(NEW|REFURBISHED)\s*\)?|\b(NEW|REFURBISHED|REFURB)\b", " ", t)
    for p in parts:
        t = t.replace(p, " ")
    t = re.sub(r"P/?N:?", " ", t)
    t = re.sub(r"\(\s*\)?|\(\s*$", " ", t)
    t = re.sub(r"[()]", " ", t)
    t = re.sub(r"\s+", " ", t).strip(" -/")
    guessed = False
    if re.search(r"\bSONOSITE\b", t):          # e.g. "L12-3 SONOSITE"
        maker, t, guessed = "SonoSite / Fujifilm", re.sub(r"\s*\bSONOSITE\b\s*", " ", t).strip(), True
    if not maker:
        probe = t
        for pat, name in MODEL_MAKERS:
            if re.search(pat, probe) or (parts and re.search(pat, " ".join(parts))):
                maker, guessed = name, True
                break
    return maker, t, part, cond, guessed


def category_for(model):
    if re.search(SYSTEM_PART, model):
        return "System part"
    if re.search(SYSTEM, model):
        return "Ultrasound system"
    if re.search(ACCESSORY, model):
        return "Accessory"
    return "Ultrasound probe"


def dom_from(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{v.year:04d}-{v.month:02d}", None
    s = clean(v)
    if not s:
        return "", None
    if re.fullmatch(r"(19|20)\d\d", s):
        return s, None
    m = re.fullmatch(r"((19|20)\d\d)-(\d\d)(-\d\d)?( 00:00:00)?", s)
    if m and 1 <= int(m.group(3)) <= 12:
        return f"{m.group(1)}-{m.group(3)}", None
    return "", s  # not a date: goes to notes


def customer_for(info):
    u = clean(info).upper()
    if not u or re.search(NOT_A_CUSTOMER, u):
        return ""
    for pat, name in CUSTOMERS:
        if re.search(pat, u):
            return name
    return ""


# --------------------------------------------------------------------------------------------
# Step 1: spreadsheet -> review workbook
# --------------------------------------------------------------------------------------------
REVIEW_COLS = [
    ("Row", 7), ("Orig. Location", 11), ("Orig. MODULE", 30), ("Orig. SERIAL", 15), ("Orig. COST", 11),
    ("Orig. DOM", 11), ("Orig. INFO", 30), ("Orig. other", 14),
    ("Tag", 14), ("Status", 14), ("Bin", 9), ("Manufacturer", 18), ("Model", 22), ("Part number", 15),
    ("Category", 17), ("Condition", 15), ("Serial", 15), ("Cost ($)", 11), ("DOM", 9), ("Customer", 28),
    ("Notes", 40), ("Check", 8), ("Why", 46),
]
EDITABLE_FROM = "Status"


def build_rows(src_path):
    wb = openpyxl.load_workbook(src_path, data_only=True)
    ws = wb["Location"]
    out = []
    for row_no, cells in enumerate(ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True), start=3):
        cells = list(cells) + [None] * 12
        loc, module, serial, cost, dom, info = cells[:6]
        other = [c for c in cells[6:12] if clean(c)]
        if not any(clean(c) for c in cells[:12]):
            continue
        why, notes = [], []
        status, bin_, loc_issue = status_and_bin(loc)
        if loc_issue:
            why.append(loc_issue)
        maker, model, part, cond, guessed = parse_module(module)
        if not clean(module):
            why.append("No MODULE")
        # The untitled column (K) names the manufacturer of a shelf space on that space's first row.
        # It is applied to the whole space after this loop (see SHELF MARKERS below).
        shelf_mark = ""
        for c in other:
            cu = clean(c).upper()
            if cu in ("GE",):
                shelf_mark = "GE Healthcare"
            elif cu.startswith("PHIL"):
                shelf_mark = "Philips"
            elif isinstance(c, (dt.datetime, dt.date)):
                notes.append(f"Other column: {c:%Y-%m-%d}")
            else:
                notes.append(f"Other column: {clean(c)}")
        if not maker and not shelf_mark:
            why.append("Manufacturer not recognized")
        if model == "19-5":
            why.append("Model 19-5 is probably SonoSite L19-5")
        if "SONOSITE" in clean(module).upper() and model.startswith("L12-3"):
            why.append("L12-3 marked SONOSITE: check the manufacturer")
        info_s = clean(info)
        iu = info_s.upper()
        if re.search(r"OPEN BOX", iu):
            cond = "New - open box"
        elif not cond and re.fullmatch(r"NEW", iu):
            cond = "New"
        elif not cond and re.search(r"\bDEMO\b", iu):
            cond = "Demo"
        cond = cond or "Used"
        # Serial
        ser = serial
        if isinstance(ser, float) and not ser.is_integer() or (isinstance(ser, float) and ser > 1e15):
            # Excel stored a long code as a number and dropped its last digits (confirmed: only
            # sold ECG accessories). Imported with no serial; the rounded number is kept in notes.
            notes.append(f"SERIAL column held a long number Excel rounded ({format(Decimal(str(ser)), 'f')}); real digits lost")
            ser = ""
        else:
            ser = clean(ser).upper()
        if not ser and not any("Excel rounded" in n for n in notes):
            notes.append("No serial in the spreadsheet")
        # Cost
        cost_v, cost_note = None, None
        if isinstance(cost, (int, float)):
            if cost > 100000:
                why.append(f"Cost {cost:,.0f} looks wrong (a serial typed in the cost column?); left blank")
            elif cost > 0:
                cost_v = round(float(cost), 2)
        elif clean(cost):
            d, rest = dom_from(cost)
            if d:
                cost_note = d
                why.append(f'Cost column held a date ({clean(cost)}); moved to DOM')
            else:
                notes.append(f"Cost column: {clean(cost)}")
        # DOM
        dom_v, dom_rest = dom_from(dom)
        if dom_rest:
            notes.append(f"DOM column: {dom_rest}")
        if not dom_v and cost_note:
            dom_v = cost_note
        m = re.search(r"DOM ((19|20)\d\d)", iu)
        if not dom_v and m:
            dom_v = m.group(1)
        # Customer
        cust = customer_for(info_s) if status in CUSTOMER_STATUSES else ""
        if status == "On loan" and not cust:
            why.append("On loan but INFO names no customer")
        if info_s:
            notes.insert(0, info_s)
        cat = category_for(model)
        out.append({
            "Row": row_no, "Orig. Location": clean(loc), "Orig. MODULE": clean(module),
            "Orig. SERIAL": clean(serial) if not isinstance(serial, float) else str(serial),
            "Orig. COST": cost if isinstance(cost, (int, float)) else clean(cost),
            "Orig. DOM": dom.strftime("%Y-%m-%d") if isinstance(dom, (dt.datetime, dt.date)) else clean(dom),
            "Orig. INFO": info_s, "Orig. other": "; ".join(clean(c) if not isinstance(c, dt.datetime) else f"{c:%Y-%m-%d}" for c in other),
            "Tag": f"CC-IMP-{row_no:04d}", "Status": status, "Bin": bin_, "Manufacturer": maker, "Model": model,
            "Part number": part, "Category": cat, "Condition": cond, "Serial": ser, "Cost ($)": cost_v,
            "DOM": dom_v, "Customer": cust, "Notes": " | ".join(notes), "_guessed": guessed,
            "Check": "CHECK" if why else "", "Why": "; ".join(why), "_mark": shelf_mark,
        })
    # SHELF MARKERS: a manufacturer written beside a shelf space's first row applies to every row
    # stored in that space (A1 covers A1 and A1/R2). A maker written in MODULE wins; a maker only
    # guessed from the model name that disagrees with the shelf is flagged for a person.
    space = lambda b: b.split("/")[0] if b else ""
    marks = {}
    for r in out:
        if r["_mark"] and space(r["Bin"]):
            marks.setdefault(space(r["Bin"]), r["_mark"])
    for r in out:
        mark = marks.get(space(r["Bin"]))
        if not mark:
            continue
        add = lambda w: (r.update(Why=(r["Why"] + "; " if r["Why"] else "") + w, Check="CHECK"))
        if not r["Manufacturer"]:
            r["Manufacturer"], r["_guessed"] = mark, False
            r["Why"] = "; ".join(w for w in r["Why"].split("; ") if w and w != "Manufacturer not recognized")
            r["Check"] = "CHECK" if r["Why"] else ""
        elif r["Manufacturer"] != mark and r["_guessed"]:
            add(f"Shelf {space(r['Bin'])} is marked {mark}, but the model looks like {r['Manufacturer']}")
        elif r["Manufacturer"] == mark:
            r["_guessed"] = False
    # Serials on more than one row: note it (not a problem - it's history).
    seen = defaultdict(list)
    for r in out:
        if r["Serial"]:
            seen[r["Serial"]].append(r["Row"])
    for r in out:
        others = [n for n in seen.get(r["Serial"], []) if n != r["Row"]]
        if others:
            note = "Same serial on row " + ", ".join(map(str, others))
            r["Notes"] = (r["Notes"] + " | " if r["Notes"] else "") + note
    return out


def write_review(rows, dst):
    wb = openpyxl.Workbook()
    head = Font(name="Arial", bold=True, color="FFFFFF")
    body = Font(name="Arial", size=10)
    orig_fill = PatternFill("solid", fgColor="EEEEEE")
    edit_fill = PatternFill("solid", fgColor="FFF7D6")
    check_fill = PatternFill("solid", fgColor="F8D7DA")
    purple = PatternFill("solid", fgColor="6A4694")
    grey_head = PatternFill("solid", fgColor="666666")

    # Rows sheet
    ws = wb.active
    ws.title = "Rows"
    names = [c for c, _ in REVIEW_COLS]
    edit_start = names.index(EDITABLE_FROM)
    for j, (name, width) in enumerate(REVIEW_COLS, start=1):
        c = ws.cell(row=1, column=j, value=name)
        c.font = head
        c.fill = grey_head if j - 1 < names.index("Tag") else purple
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = width
    for i, r in enumerate(rows, start=2):
        for j, name in enumerate(names, start=1):
            c = ws.cell(row=i, column=j, value=r[name] if r[name] not in ("", None) else None)
            c.font = body
            if j - 1 < names.index("Tag"):
                c.fill = orig_fill
            elif j - 1 >= edit_start and name not in ("Check", "Why"):
                c.fill = edit_fill
            if name == "Check" and r["Check"]:
                c.fill = check_fill
                c.font = Font(name="Arial", size=10, bold=True, color="9C0006")
            if name == "Cost ($)":
                c.number_format = "$#,##0.00"
            if name in ("Orig. SERIAL", "Serial", "Part number", "DOM", "Orig. DOM"):
                c.number_format = "@"
    ws.freeze_panes = "J2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(names))}{len(rows) + 1}"
    ws.row_dimensions[1].height = 30
    last = len(rows) + 1
    col = {n: get_column_letter(names.index(n) + 1) for n in names}

    # Customers sheet
    wc = wb.create_sheet("Customers")
    spell = defaultdict(set)
    for r in rows:
        if r["Customer"]:
            spell[r["Customer"]].add(r["Orig. INFO"])
    wc.append(["Customer (as it will appear in the app)", "Spellings found in INFO", "Records linked"])
    for j in range(1, 4):
        wc.cell(row=1, column=j).font = head
        wc.cell(row=1, column=j).fill = purple
    for k, name in enumerate(sorted(spell), start=2):
        wc.cell(row=k, column=1, value=name).font = body
        wc.cell(row=k, column=2, value=" / ".join(sorted(spell[name]))).font = body
        wc.cell(row=k, column=3, value=f"=COUNTIF(Rows!${col['Customer']}$2:${col['Customer']}${last},A{k})").font = body
    wc.column_dimensions["A"].width = 40
    wc.column_dimensions["B"].width = 80
    wc.column_dimensions["C"].width = 16
    wc.freeze_panes = "A2"

    # Models sheet (replaces the old Count sheet: units by model, on hand vs sold)
    wm = wb.create_sheet("Models")
    combos = sorted({(r["Manufacturer"], r["Model"], r["Category"]) for r in rows}, key=lambda x: (x[0] or "~", x[1]))
    wm.append(["Manufacturer", "Model", "Category", "All records", "In stock", "On loan", "Sold / shipped"])
    for j in range(1, 8):
        wm.cell(row=1, column=j).font = head
        wm.cell(row=1, column=j).fill = purple
    M, D, S = col["Manufacturer"], col["Model"], col["Status"]
    for k, (mk, md, ct) in enumerate(combos, start=2):
        wm.cell(row=k, column=1, value=mk or "(not recognized)").font = body
        wm.cell(row=k, column=2, value=md).font = body
        wm.cell(row=k, column=3, value=ct).font = body
        mref = f'""' if not mk else f"A{k}"
        base = f"Rows!${M}$2:${M}${last},{mref},Rows!${D}$2:${D}${last},B{k}"
        wm.cell(row=k, column=4, value=f"=COUNTIFS({base})").font = body
        for jj, st in ((5, "In stock"), (6, "On loan"), (7, "Sold / shipped")):
            wm.cell(row=k, column=jj, value=f'=COUNTIFS({base},Rows!${S}$2:${S}${last},"{st}")').font = body
    for letter, w in zip("ABCDEFG", (22, 30, 18, 12, 10, 10, 14)):
        wm.column_dimensions[letter].width = w
    wm.freeze_panes = "A2"
    wm.auto_filter.ref = f"A1:G{len(combos) + 1}"

    # How to review sheet (first)
    wh = wb.create_sheet("How to review", 0)
    lines = [
        ("CC Medical Inventory - import review", Font(name="Arial", size=14, bold=True, color="6A4694")),
        ("Each row of Inventory_Ultrasound.xlsx (Location sheet) becomes one record in the app. Grey columns are the original; yellow columns are what the app will store.", body),
        ("1. On the Rows sheet, filter Check = CHECK and fix those rows first. The Why column says what to look at.", body),
        ("2. Then skim the Models sheet: one line per model with its manufacturer. A wrong manufacturer there means fixing it on the Rows sheet (filter by Model).", body),
        ("3. On the Customers sheet, check the customer names made from the INFO notes. To rename one, change it in the Rows sheet's Customer column.", body),
        ("4. Edit only the yellow columns. Leave Tag and Row as they are; they tie each record back to its spreadsheet row.", body),
        ("Example: Row 9 reads Location A1, MODULE M5SC-D, SERIAL 518112YP8, so the app stores Status In stock, Bin A1, Manufacturer GE Healthcare, Model M5SC-D, Category Ultrasound probe.", body),
        ("Status must be one of: In stock, Pending, On loan, Out on rental, Sold / shipped. Cost is dollars per unit. DOM is YYYY-MM or YYYY.", body),
        ("Shelf manufacturers: the maker written in the untitled column on a shelf space's first row (GE on A1-A3, Philips on A4, B1, B2) is applied to every row in that space. A maker written in MODULE (ACUSON, SIEMENS, SONOSITE...) still wins; rows where the model suggests a different maker than the shelf are marked CHECK.", body),
        ("", body),
        ("Counts", Font(name="Arial", size=12, bold=True)),
    ]
    for i, (text, f) in enumerate(lines, start=1):
        c = wh.cell(row=i, column=1, value=text)
        c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")
    stats = [
        ("Rows to import", f"=COUNTA(Rows!A2:A{last})"),
        ("Rows marked CHECK", f'=COUNTIF(Rows!{col["Check"]}2:{col["Check"]}{last},"CHECK")'),
        ("In stock", f'=COUNTIF(Rows!{S}2:{S}{last},"In stock")'),
        ("Pending", f'=COUNTIF(Rows!{S}2:{S}{last},"Pending")'),
        ("On loan", f'=COUNTIF(Rows!{S}2:{S}{last},"On loan")'),
        ("Out on rental", f'=COUNTIF(Rows!{S}2:{S}{last},"Out on rental")'),
        ("Sold / shipped", f'=COUNTIF(Rows!{S}2:{S}{last},"Sold / shipped")'),
        ("Manufacturer not recognized", f'=COUNTBLANK(Rows!{M}2:{M}{last})'),
        ("Records with a cost", f'=COUNT(Rows!{col["Cost ($)"]}2:{col["Cost ($)"]}{last})'),
        ("Recorded cost of units in stock, pending or on loan ($)",
         f'=SUMIFS(Rows!{col["Cost ($)"]}2:{col["Cost ($)"]}{last},Rows!{S}2:{S}{last},"In stock")+SUMIFS(Rows!{col["Cost ($)"]}2:{col["Cost ($)"]}{last},Rows!{S}2:{S}{last},"Pending")+SUMIFS(Rows!{col["Cost ($)"]}2:{col["Cost ($)"]}{last},Rows!{S}2:{S}{last},"On loan")'),
        ("Customers created", f"=COUNTA(Customers!A2:A{len(spell) + 1})"),
    ]
    r0 = len(lines) + 1
    for k, (label, formula) in enumerate(stats):
        wh.cell(row=r0 + k, column=1, value=label).font = body
        c = wh.cell(row=r0 + k, column=2, value=formula)
        c.font = body
        if "$" in label:
            c.number_format = "$#,##0.00"
    wh.column_dimensions["A"].width = 110
    wh.column_dimensions["B"].width = 16
    wb.save(dst)


# --------------------------------------------------------------------------------------------
# Step 2: reviewed workbook -> SQL
# --------------------------------------------------------------------------------------------
def q(v):
    if v is None or v == "":
        return "''"
    return "'" + str(v).replace("'", "''") + "'"


def qn(v):
    return "NULL" if v in (None, "") else repr(round(float(v), 2))


def read_review(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Rows"]
    hdr = [clean(c.value) for c in ws[1]]
    need = ["Row", "Tag", "Status", "Bin", "Manufacturer", "Model", "Part number", "Category", "Condition",
            "Serial", "Cost ($)", "DOM", "Customer", "Notes"]
    missing = [n for n in need if n not in hdr]
    if missing:
        sys.exit(f"Review file is missing columns: {', '.join(missing)}")
    ix = {n: hdr.index(n) for n in hdr}
    rows, problems = [], []
    for cells in ws.iter_rows(min_row=2, values_only=True):
        if cells[ix["Tag"]] in (None, ""):
            continue
        r = {n: cells[ix[n]] for n in need}
        r["Row"] = int(r["Row"])
        for n in need:
            if n not in ("Row", "Cost ($)"):
                r[n] = clean(r[n])
        r["Serial"] = r["Serial"].upper()
        if r["Status"] not in STOCK_STATUSES:
            problems.append(f"Row {r['Row']}: Status '{r['Status']}' is not one of {', '.join(STOCK_STATUSES)}")
        if not r["Category"]:
            problems.append(f"Row {r['Row']}: Category is blank")
        if r["DOM"] and not re.fullmatch(r"(19|20)\d\d(-(0[1-9]|1[0-2]))?", r["DOM"]):
            problems.append(f"Row {r['Row']}: DOM '{r['DOM']}' should look like 2024-08 or 2024")
        if r["Cost ($)"] not in (None, ""):
            try:
                r["Cost ($)"] = float(str(r["Cost ($)"]).replace("$", "").replace(",", ""))
            except ValueError:
                problems.append(f"Row {r['Row']}: Cost '{r['Cost ($)']}' is not a number")
        rows.append(r)
    if problems:
        sys.exit("Fix these in the review file, then run again:\n  " + "\n  ".join(problems[:50]))
    return rows


def write_sql(rows, dst, when=None):
    t = when or dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")
    by = "Spreadsheet import"
    out = [
        "-- CC Medical Inventory: import of Inventory_Ultrasound.xlsx",
        f"-- Generated {t}. {len(rows)} records.",
        "-- Load with: npx wrangler d1 execute cc-inventory --remote --file=import.sql",
        "-- Safe to run more than once: anything already imported is skipped.",
        "",
    ]
    cust_id = {}
    for r in rows:
        if r["Customer"] and r["Customer"] not in cust_id:
            cust_id[r["Customer"]] = "c_" + hashlib.sha1(r["Customer"].lower().encode()).hexdigest()[:16]
    out.append("-- Customers")
    for name, cid in sorted(cust_id.items()):
        out.append(f"INSERT INTO customers (id, name, type, facility, phone, email, address, notes, created_at, updated_at) "
                   f"SELECT {q(cid)}, {q(name)}, '', '', '', '', '', 'Added by the spreadsheet import', {q(t)}, {q(t)} "
                   f"WHERE NOT EXISTS (SELECT 1 FROM customers WHERE name = {q(name)} COLLATE NOCASE);")
    out += ["", "-- Records and their history"]
    for r in rows:
        name = " ".join(x for x in (r["Manufacturer"], r["Model"]) if x) or r["Category"]
        cid = f"(SELECT id FROM customers WHERE name = {q(r['Customer'])} COLLATE NOCASE LIMIT 1)" if r["Customer"] else "NULL"
        out.append(
            "INSERT OR IGNORE INTO items (id, kind, name, category, cond, manufacturer, model, part_number, serial, ref, cost, dom, source, "
            "qty, location, customer_id, status, problems, notes, received_by, received_at, updated_at) VALUES ("
            f"{q(r['Tag'])}, 'stock', {q(name)}, {q(r['Category'])}, {q(r['Condition'])}, {q(r['Manufacturer'])}, {q(r['Model'])}, "
            f"{q(r['Part number'])}, {q(r['Serial'])}, '', {qn(r['Cost ($)'])}, {q(r['DOM'])}, {q('import:' + str(r['Row']))}, "
            f"1, {q(r['Bin'])}, {cid}, {q(r['Status'])}, '', {q(r['Notes'])}, {q(by)}, {q(t)}, {q(t)});")
        what = f"Imported from Inventory_Ultrasound.xlsx, row {r['Row']} ({r['Status']}{', bin ' + r['Bin'] if r['Bin'] else ''})"
        out.append(f"INSERT INTO history (item_id, at, by, what) SELECT {q(r['Tag'])}, {q(t)}, {q(by)}, {q(what)} "
                   f"WHERE NOT EXISTS (SELECT 1 FROM history WHERE item_id = {q(r['Tag'])} AND what = {q(what)});")
        if r["Serial"]:
            out.append(f"INSERT INTO serial_events (serial, at, by, item_id, kind, what) SELECT {q(r['Serial'])}, {q(t)}, {q(by)}, {q(r['Tag'])}, 'stock', {q(what)} "
                       f"WHERE NOT EXISTS (SELECT 1 FROM serial_events WHERE item_id = {q(r['Tag'])} AND what = {q(what)});")
    # Serial registry: one row per unit; its latest record = the one still open, else the last row.
    out += ["", "-- Product History (serial registry)"]
    by_serial = defaultdict(list)
    for r in rows:
        if r["Serial"]:
            by_serial[r["Serial"]].append(r)
    for s, rs in sorted(by_serial.items()):
        open_rs = [r for r in rs if r["Status"] not in CLOSED]
        latest = (open_rs or rs)[-1]
        out.append(
            "INSERT OR IGNORE INTO serials (serial, manufacturer, model, part_number, category, dom, times_received, first_seen, last_seen, last_item_id) VALUES ("
            f"{q(s)}, {q(latest['Manufacturer'])}, {q(latest['Model'])}, {q(latest['Part number'])}, {q(latest['Category'])}, "
            f"{q(next((r['DOM'] for r in rs if r['DOM']), ''))}, {len(rs)}, {q(t)}, {q(t)}, {q(latest['Tag'])});")
    # Buttons
    out += ["", "-- Buttons: manufacturers, models, categories, conditions, bins"]
    opts = set()
    for r in rows:
        opts.add(("manufacturer", r["Manufacturer"], ""))
        opts.add(("model", r["Model"], r["Manufacturer"]))
        opts.add(("category", r["Category"], ""))
        opts.add(("condition", r["Condition"], ""))
        for b in r["Bin"].split("/"):
            opts.add(("location", b.strip(), ""))
    for kind, value, parent in sorted(o for o in opts if o[1]):
        out.append(f"INSERT OR IGNORE INTO options (kind, value, parent, created_at) VALUES ({q(kind)}, {q(value)}, {q(parent)}, {q(t)});")
    out.append("")
    with open(dst, "w", encoding="utf-8") as f:
        f.write("\n".join(out))
    return {"records": len(rows), "customers": len(cust_id), "serials": len(by_serial), "options": len([o for o in opts if o[1]])}


def write_json(rows, dst):
    """The same import as write_sql, as a file an admin uploads on the app's Admin page."""
    by_serial = defaultdict(list)
    for r in rows:
        if r["Serial"]:
            by_serial[r["Serial"]].append(r)
    serials = []
    for s, rs in sorted(by_serial.items()):
        open_rs = [r for r in rs if r["Status"] not in CLOSED]
        latest = (open_rs or rs)[-1]
        serials.append({"serial": s, "manufacturer": latest["Manufacturer"], "model": latest["Model"],
                        "part_number": latest["Part number"], "category": latest["Category"],
                        "dom": next((r["DOM"] for r in rs if r["DOM"]), ""), "times": len(rs), "last_item_id": latest["Tag"]})
    opts = set()
    for r in rows:
        opts.add(("manufacturer", r["Manufacturer"], ""))
        opts.add(("model", r["Model"], r["Manufacturer"]))
        opts.add(("category", r["Category"], ""))
        opts.add(("condition", r["Condition"], ""))
        for b in r["Bin"].split("/"):
            opts.add(("location", b.strip(), ""))
    data = {
        "format": "cc-medical-import-1",
        "generated": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z"),
        "customers": sorted({r["Customer"] for r in rows if r["Customer"]}),
        "items": [{"tag": r["Tag"], "name": " ".join(x for x in (r["Manufacturer"], r["Model"]) if x) or r["Category"],
                   "category": r["Category"], "cond": r["Condition"], "manufacturer": r["Manufacturer"], "model": r["Model"],
                   "part_number": r["Part number"], "serial": r["Serial"], "cost": r["Cost ($)"], "dom": r["DOM"], "row": r["Row"],
                   "location": r["Bin"], "customer": r["Customer"], "status": r["Status"], "notes": r["Notes"]} for r in rows],
        "serials": serials,
        "options": [list(o) for o in sorted(o for o in opts if o[1])],
    }
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"), default=str)
    return {"records": len(data["items"]), "customers": len(data["customers"]), "serials": len(serials), "options": len(data["options"])}


CSV_COLUMNS = ["Tag", "Manufacturer", "Model", "Category", "Condition", "Part number", "Serial", "Cost",
               "Date of manufacture", "Bin", "Status", "Customer", "Notes", "Spreadsheet row"]


def write_csv(rows, dst):
    """The same import as a CSV that opens in Excel. The app's Admin page imports it directly."""
    import csv
    with open(dst, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(CSV_COLUMNS)
        for r in rows:
            cost = r["Cost ($)"]
            w.writerow([r["Tag"], r["Manufacturer"], r["Model"], r["Category"], r["Condition"], r["Part number"], r["Serial"],
                        "" if cost in (None, "") else (int(cost) if float(cost).is_integer() else cost),
                        r["DOM"], r["Bin"], r["Status"], r["Customer"], r["Notes"], r["Row"]])
    return {"records": len(rows)}


if __name__ == "__main__":
    if len(sys.argv) != 4 or sys.argv[1] not in ("review", "sql", "json", "csv"):
        sys.exit(__doc__)
    if sys.argv[1] == "review":
        rows = build_rows(sys.argv[2])
        write_review(rows, sys.argv[3])
        flagged = sum(1 for r in rows if r["Check"])
        print(f"Wrote {sys.argv[3]}: {len(rows)} rows, {flagged} marked CHECK.")
    elif sys.argv[1] == "csv":
        stats = write_csv(read_review(sys.argv[2]), sys.argv[3])
        print(f"Wrote {sys.argv[3]}: {stats}")
    elif sys.argv[1] == "json":
        stats = write_json(read_review(sys.argv[2]), sys.argv[3])
        print(f"Wrote {sys.argv[3]}: {stats}")
    else:
        stats = write_sql(read_review(sys.argv[2]), sys.argv[3])
        print(f"Wrote {sys.argv[3]}: {stats}")
